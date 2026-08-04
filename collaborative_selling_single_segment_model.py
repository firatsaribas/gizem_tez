"""
Collaborative Selling Model — Single-Segment Threshold Revenue Structure
========================================================================

Updated model logic
-------------------
This version implements the revised mathematical model in which each active
collaborative group and each individual seller selects exactly one revenue
segment.

A selling entity can select segment l only if its total quantity lies within
the interval [A_l, B_l]. Therefore:
    - small quantities cannot access premium segments;
    - large quantities cannot simultaneously access lower segments;
    - each group's or individual's entire quantity is evaluated under one
      selected segment.

This version replaces the previous "access gate" interpretation of z variables
with a "selected segment" interpretation.
"""

import math
from docplex.mp.model import Model

# =============================================================================
# 1. INSTANCE DATA
# =============================================================================

producers = list(range(1, 11))
segments = [0, 1, 2]
G = list(range(len(producers)))  # potential group indices

district = {
    1: "Torbalı, İzmir",
    2: "Selçuk, İzmir",
    3: "Urla, İzmir",
    4: "Foça, İzmir",
    5: "Ödemiş, İzmir",
    6: "Nazilli, Aydın",
    7: "Söke, Aydın",
    8: "Akhisar, Manisa",
    9: "Kırkağaç, Manisa",
    10: "Milas, Muğla",
}

lat = {
    1: 38.163, 2: 37.956, 3: 38.321, 4: 38.673, 5: 38.226,
    6: 37.913, 7: 37.748, 8: 38.919, 9: 39.101, 10: 37.318,
}

lon = {
    1: 27.359, 2: 27.369, 3: 26.765, 4: 26.755, 5: 27.971,
    6: 28.317, 7: 27.413, 8: 27.837, 9: 27.668, 10: 27.787,
}

q = {
    1: 12, 2: 7, 3: 15, 4: 10, 5: 5,
    6: 18, 7: 6, 8: 14, 9: 10, 10: 9,
}

# ── Distances via Haversine ───────────────────────────────────────────────────
def haversine(lat1, lon1, lat2, lon2):
    R = 6371.0
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    a = (
        math.sin(math.radians(lat2 - lat1) / 2) ** 2
        + math.cos(phi1) * math.cos(phi2)
        * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    )
    return 2.0 * R * math.asin(math.sqrt(a))


d = {
    (i, j): haversine(lat[i], lon[i], lat[j], lon[j])
    for i in producers for j in producers if i != j
}

D_max = 75.0

infeasible_pairs = [
    (i, j)
    for i in producers for j in producers
    if i < j and d[(i, j)] > D_max
]

# ── Segment parameters ────────────────────────────────────────────────────────
# A[l]: lower threshold of segment l
# B[l]: upper threshold/capacity of segment l
# p[l]: unit revenue of segment l
A = {0: 0, 1: 5, 2: 25}
B = {0: 5, 1: 25, 2: 30}
p = {0: 130, 1: 175, 2: 220}

segment_label = {
    0: "Base segment",
    1: "Medium-premium segment",
    2: "High-premium segment",
}

# Big-M should be at least as large as the maximum possible group quantity.
M = sum(q.values())

# =============================================================================
# 2. INSTANCE SUMMARY
# =============================================================================

sep = "=" * 72
print(sep)
print("BASE CASE — Single-Segment Threshold Revenue Model")
print("N=10 producers | |G|=10 potential groups | L=3 revenue segments")
print(sep)
print(f"Total supply : {sum(q.values())} t/season")
print(f"D_max        : {D_max} km")
print(f"Big-M        : {M}")
print()

print(f"{'#':>2}  {'District':<24} {'Lat':>7}  {'Lon':>7}  {'q_i (t)':>8}")
print("─" * 58)
for i in producers:
    print(f"{i:>2}  {district[i]:<24} {lat[i]:>7.3f}  {lon[i]:>7.3f}  {q[i]:>8}")

print()
print(f"{'Seg':>3}  {'Description':<28} {'A_l':>5}  {'B_l':>5}  {'p_l':>7}")
print("─" * 58)
for l in segments:
    print(f"{l:>3}  {segment_label[l]:<28} {A[l]:>5}  {B[l]:>5}  {p[l]:>7}")

print()
print(f"Infeasible pairs (d > {D_max} km): {len(infeasible_pairs)} of 45")
print(sep)

# =============================================================================
# 3. MODEL
# =============================================================================

mdl = Model(name="Collab_Selling_Single_Segment_Threshold")
mdl.parameters.mip.tolerances.mipgap = 1e-6
mdl.parameters.timelimit = 3600

# ── Decision variables ────────────────────────────────────────────────────────
x = {
    (i, g): mdl.binary_var(name=f"x_{i}_{g}")
    for i in producers for g in G
}

s = {
    i: mdl.binary_var(name=f"s_{i}")
    for i in producers
}

y = {
    g: mdl.binary_var(name=f"y_{g}")
    for g in G
}

Q = {
    g: mdl.continuous_var(lb=0, name=f"Q_{g}")
    for g in G
}

# z variables now mean SELECTED segment, not merely accessible segment.
zG = {
    (g, l): mdl.binary_var(name=f"zG_{g}_{l}")
    for g in G for l in segments
}

zI = {
    (i, l): mdl.binary_var(name=f"zI_{i}_{l}")
    for i in producers for l in segments
}

rG = {
    (g, l): mdl.continuous_var(lb=0, name=f"rG_{g}_{l}")
    for g in G for l in segments
}

rI = {
    (i, l): mdl.continuous_var(lb=0, name=f"rI_{i}_{l}")
    for i in producers for l in segments
}

# ── Objective ─────────────────────────────────────────────────────────────────
mdl.maximize(
    mdl.sum(p[l] * rG[(g, l)] for g in G for l in segments)
    + mdl.sum(p[l] * rI[(i, l)] for i in producers for l in segments)
)

# =============================================================================
# 4. CONSTRAINTS
# =============================================================================

# (1) Each producer either joins exactly one group or sells individually.
for i in producers:
    mdl.add_constraint(
        mdl.sum(x[(i, g)] for g in G) + s[i] == 1,
        ctname=f"C_assign_i{i}",
    )

# (2) A producer can join only an active group.
for i in producers:
    for g in G:
        mdl.add_constraint(
            x[(i, g)] <= y[g],
            ctname=f"C_active_i{i}_g{g}",
        )

# (3) Active group must include at least two producers.
for g in G:
    mdl.add_constraint(
        mdl.sum(x[(i, g)] for i in producers) >= 2 * y[g],
        ctname=f"C_minsize_g{g}",
    )

# (4) If group is inactive, no producer can be assigned to it.
# This is redundant with x_ig <= y_g, but kept here to match the paper model.
for g in G:
    mdl.add_constraint(
        mdl.sum(x[(i, g)] for i in producers) <= len(producers) * y[g],
        ctname=f"C_inactive_g{g}",
    )

# (5) Group quantity.
for g in G:
    mdl.add_constraint(
        Q[g] == mdl.sum(q[i] * x[(i, g)] for i in producers),
        ctname=f"C_qty_g{g}",
    )

# (6) Each active group selects exactly one segment.
for g in G:
    mdl.add_constraint(
        mdl.sum(zG[(g, l)] for l in segments) == y[g],
        ctname=f"C_one_segment_G_g{g}",
    )

# (7) Each individual seller selects exactly one segment.
for i in producers:
    mdl.add_constraint(
        mdl.sum(zI[(i, l)] for l in segments) == s[i],
        ctname=f"C_one_segment_I_i{i}",
    )

# (8) Group lower and upper segment interval constraints.
for g in G:
    for l in segments:
        mdl.add_constraint(
            Q[g] >= A[l] * zG[(g, l)],
            ctname=f"C_lower_G_g{g}_l{l}",
        )
        mdl.add_constraint(
            Q[g] <= B[l] * zG[(g, l)] + M * (1 - zG[(g, l)]),
            ctname=f"C_upper_G_g{g}_l{l}",
        )

# (9) Individual lower and upper segment interval constraints.
for i in producers:
    for l in segments:
        mdl.add_constraint(
            q[i] * s[i] >= A[l] * zI[(i, l)],
            ctname=f"C_lower_I_i{i}_l{l}",
        )
        mdl.add_constraint(
            q[i] * s[i] <= B[l] * zI[(i, l)] + M * (1 - zI[(i, l)]),
            ctname=f"C_upper_I_i{i}_l{l}",
        )

# (10) Group quantity can be allocated only to the selected segment.
for g in G:
    for l in segments:
        mdl.add_constraint(
            rG[(g, l)] <= B[l] * zG[(g, l)],
            ctname=f"C_rG_cap_g{g}_l{l}",
        )

# (11) Full group quantity must be allocated.
for g in G:
    mdl.add_constraint(
        mdl.sum(rG[(g, l)] for l in segments) == Q[g],
        ctname=f"C_alloc_G_g{g}",
    )

# (12) Individual quantity can be allocated only to the selected segment.
for i in producers:
    for l in segments:
        mdl.add_constraint(
            rI[(i, l)] <= B[l] * zI[(i, l)],
            ctname=f"C_rI_cap_i{i}_l{l}",
        )

# (13) Full individual quantity must be allocated.
for i in producers:
    mdl.add_constraint(
        mdl.sum(rI[(i, l)] for l in segments) == q[i] * s[i],
        ctname=f"C_alloc_I_i{i}",
    )

# (14) Geographical incompatibility.
for g in G:
    for (i, j) in infeasible_pairs:
        mdl.add_constraint(
            x[(i, g)] + x[(j, g)] <= 1,
            ctname=f"C_geo_i{i}_j{j}_g{g}",
        )

# (15) Symmetry breaking: active groups occupy the lowest indices first.
for g in range(len(G) - 1):
    mdl.add_constraint(
        y[g] >= y[g + 1],
        ctname=f"C_sym_g{g}",
    )

# =============================================================================
# 5. MODEL SIZE
# =============================================================================

n_bin = len(x) + len(s) + len(y) + len(zG) + len(zI)
n_cont = len(Q) + len(rG) + len(rI)

print("\nMODEL SIZE")
print(f"Binary variables     : {n_bin}")
print(f"Continuous variables : {n_cont}")
print(f"Total variables      : {n_bin + n_cont}")
print(f"Constraints          : {mdl.number_of_constraints}")
print()

# =============================================================================
# 6. SOLVE
# =============================================================================

print("Solving with IBM CPLEX ...")
sol = mdl.solve(log_output=True)
print()

if sol is None:
    print("ERROR: No feasible solution found.")
else:
    obj_val = sol.objective_value

    print(sep)
    print("SOLUTION REPORT")
    print(sep)
    print(f"Solver status   : {mdl.solve_details.status}")
    print(f"MIP gap         : {mdl.solve_details.mip_relative_gap:.2e}")
    print(f"Optimal revenue : ${obj_val:,.2f}")
    print()

    # Individual sellers
    ind = [i for i in producers if sol.get_value(s[i]) > 0.5]
    ind_rev = 0.0

    print(f"INDIVIDUAL SELLERS ({len(ind)} producers)")
    for i in ind:
        selected_segment = next(
            l for l in segments if sol.get_value(zI[(i, l)]) > 0.5
        )
        rev_i = sum(p[l] * sol.get_value(rI[(i, l)]) for l in segments)
        ind_rev += rev_i
        print(
            f"  Producer {i}: {district[i]:<24} "
            f"q={q[i]:>2}t | selected seg={selected_segment} "
            f"({segment_label[selected_segment]}) | revenue=${rev_i:,.2f}"
        )

    # Collaborative groups
    act = [g for g in G if sol.get_value(y[g]) > 0.5]
    grp_rev = 0.0

    print()
    print(f"COLLABORATIVE GROUPS ({len(act)} active)")
    for g in act:
        mems = [i for i in producers if sol.get_value(x[(i, g)]) > 0.5]
        tot_qty = sum(q[i] for i in mems)
        selected_segment = next(
            l for l in segments if sol.get_value(zG[(g, l)]) > 0.5
        )
        rev_g = sum(p[l] * sol.get_value(rG[(g, l)]) for l in segments)
        grp_rev += rev_g

        print(f"  Group {g}: members {mems} -> {[district[i] for i in mems]}")
        print(
            f"    Q={tot_qty}t | selected seg={selected_segment} "
            f"({segment_label[selected_segment]}) | revenue=${rev_g:,.2f}"
        )

    print()
    print(f"Revenue — individual : ${ind_rev:,.2f}")
    print(f"Revenue — groups     : ${grp_rev:,.2f}")
    print(f"TOTAL OPTIMAL REVENUE: ${obj_val:,.2f}")
    print(sep)

    # =============================================================================
    # 7. VALIDATION
    # =============================================================================

    print("\nVALIDATION")

    # (a) All producers assigned exactly once
    assigned = set(ind)
    for g in act:
        assigned |= {i for i in producers if sol.get_value(x[(i, g)]) > 0.5}
    ok_assign = assigned == set(producers)
    print(f"[{'OK' if ok_assign else 'FAIL'}] All producers assigned")

    # (b) Geographical constraints
    geo_ok = all(
        d[(mems[a], mems[b])] <= D_max
        for g in act
        for mems in [[i for i in producers if sol.get_value(x[(i, g)]) > 0.5]]
        for a in range(len(mems))
        for b in range(a + 1, len(mems))
    )
    print(f"[{'OK' if geo_ok else 'FAIL'}] Geographical constraints satisfied")

    # (c) One selected segment for each active entity
    one_seg_group_ok = all(
        abs(sum(sol.get_value(zG[(g, l)]) for l in segments) - 1) < 1e-6
        for g in act
    )

    one_seg_ind_ok = all(
        abs(sum(sol.get_value(zI[(i, l)]) for l in segments) - 1) < 1e-6
        for i in ind
    )

    print(f"[{'OK' if one_seg_group_ok else 'FAIL'}] Each active group selects one segment")
    print(f"[{'OK' if one_seg_ind_ok else 'FAIL'}] Each individual seller selects one segment")

    # (d) Selected segment interval validation
    interval_group_ok = True
    for g in act:
        selected_segment = next(
            l for l in segments if sol.get_value(zG[(g, l)]) > 0.5
        )
        Q_val = sol.get_value(Q[g])
        if not (A[selected_segment] - 1e-6 <= Q_val <= B[selected_segment] + 1e-6):
            interval_group_ok = False
            break

    interval_ind_ok = True
    for i in ind:
        selected_segment = next(
            l for l in segments if sol.get_value(zI[(i, l)]) > 0.5
        )
        q_val = q[i]
        if not (A[selected_segment] - 1e-6 <= q_val <= B[selected_segment] + 1e-6):
            interval_ind_ok = False
            break

    print(f"[{'OK' if interval_group_ok else 'FAIL'}] Group quantities lie within selected segment intervals")
    print(f"[{'OK' if interval_ind_ok else 'FAIL'}] Individual quantities lie within selected segment intervals")

    # (e) Objective cross-check
    recomputed_obj = (
        sum(p[l] * sol.get_value(rG[(g, l)]) for g in G for l in segments)
        + sum(p[l] * sol.get_value(rI[(i, l)]) for i in producers for l in segments)
    )

    obj_ok = abs(recomputed_obj - obj_val) < 1e-4
    print(
        f"[{'OK' if obj_ok else 'FAIL'}] Recomputed revenue "
        f"${recomputed_obj:,.2f} matches solver objective ${obj_val:,.2f}"
    )

# =============================================================================
# 8. EXPORT TO EXCEL
# =============================================================================

def write_results_to_excel(sol, filepath="collaborative_selling_single_segment_results.xlsx"):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()

    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    ctr = Alignment(horizontal="center")
    thin = Side(style="thin")
    full_bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header_row(ws, row=1):
        for cell in ws[row]:
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = ctr
            cell.border = full_bdr

    # Sheet 1: Revenue Summary
    ws0 = wb.active
    ws0.title = "Revenue Summary"
    ws0.append([
        "Entity",
        "Type",
        "Members",
        "Districts",
        "Quantity (t)",
        "Selected Segment",
        "Segment Label",
        "Unit Revenue",
        "Total Revenue",
    ])
    style_header_row(ws0)

    ind_list = [i for i in producers if sol.get_value(s[i]) > 0.5]
    act_list = [g for g in G if sol.get_value(y[g]) > 0.5]

    for i in ind_list:
        selected_segment = next(
            l for l in segments if sol.get_value(zI[(i, l)]) > 0.5
        )
        total_rev = round(
            sum(p[l] * sol.get_value(rI[(i, l)]) for l in segments), 2
        )
        ws0.append([
            f"Producer {i}",
            "Individual",
            str(i),
            district[i],
            q[i],
            selected_segment,
            segment_label[selected_segment],
            p[selected_segment],
            total_rev,
        ])

    for g in act_list:
        mems = [i for i in producers if sol.get_value(x[(i, g)]) > 0.5]
        selected_segment = next(
            l for l in segments if sol.get_value(zG[(g, l)]) > 0.5
        )
        total_qty = sum(q[i] for i in mems)
        total_rev = round(
            sum(p[l] * sol.get_value(rG[(g, l)]) for l in segments), 2
        )
        ws0.append([
            f"Group {g}",
            "Collaborative",
            str(mems),
            "; ".join(district[i] for i in mems),
            total_qty,
            selected_segment,
            segment_label[selected_segment],
            p[selected_segment],
            total_rev,
        ])

    ws0.append([])
    ws0.append(["TOTAL", "", "", "", "", "", "", "", round(sol.objective_value, 2)])
    for cell in ws0[ws0.max_row]:
        cell.font = Font(bold=True)

    for col in ["A", "B", "C", "D", "G"]:
        ws0.column_dimensions[col].width = 26

    # Sheet 2: x_ig
    ws1 = wb.create_sheet("x_ig")
    ws1.append(["i \\ g"] + list(G))
    style_header_row(ws1)
    for i in producers:
        ws1.append([i] + [round(sol.get_value(x[(i, g)])) for g in G])

    # Sheet 3: s_i
    ws2 = wb.create_sheet("s_i")
    ws2.append(["i", "s_i"])
    style_header_row(ws2)
    for i in producers:
        ws2.append([i, round(sol.get_value(s[i]))])

    # Sheet 4: y_g
    ws3 = wb.create_sheet("y_g")
    ws3.append(["g", "y_g"])
    style_header_row(ws3)
    for g in G:
        ws3.append([g, round(sol.get_value(y[g]))])

    # Sheet 5: Q_g
    ws4 = wb.create_sheet("Q_g")
    ws4.append(["g", "Q_g"])
    style_header_row(ws4)
    for g in G:
        ws4.append([g, round(sol.get_value(Q[g]), 4)])

    # Sheet 6: zG_gl
    ws5 = wb.create_sheet("zG_gl")
    ws5.append(["g \\ l (selected segment)"] + [f"seg {l}" for l in segments])
    style_header_row(ws5)
    for g in G:
        ws5.append([g] + [round(sol.get_value(zG[(g, l)])) for l in segments])

    # Sheet 7: zI_il
    ws6 = wb.create_sheet("zI_il")
    ws6.append(["i \\ l (selected segment)"] + [f"seg {l}" for l in segments])
    style_header_row(ws6)
    for i in producers:
        ws6.append([i] + [round(sol.get_value(zI[(i, l)])) for l in segments])

    # Sheet 8: rG_gl
    ws7 = wb.create_sheet("rG_gl")
    ws7.append(["g \\ l (quantity allocated)"] + [f"seg {l}" for l in segments])
    style_header_row(ws7)
    for g in G:
        ws7.append([g] + [round(sol.get_value(rG[(g, l)]), 4) for l in segments])

    # Sheet 9: rI_il
    ws8 = wb.create_sheet("rI_il")
    ws8.append(["i \\ l (quantity allocated)"] + [f"seg {l}" for l in segments])
    style_header_row(ws8)
    for i in producers:
        ws8.append([i] + [round(sol.get_value(rI[(i, l)]), 4) for l in segments])

    wb.save(filepath)
    print(f"\nResults saved to: {filepath}")
    return filepath

if sol is not None:
    write_results_to_excel(sol, "ccollaborative_selling_single_segment_results.xlsx")
